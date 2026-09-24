import io
from openpyxl import Workbook
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import app

def make_xlsx():
    wb=Workbook(); ws=wb.active; ws.title="فاتورة"
    ws.append(["شركة تجريبية","فاتورة 100"])
    ws.append(["رقم الفاتورة","100","التاريخ","2026-09-20"])
    ws.append(["البيان","الكمية","الوحدة","سعر الوحدة","الضريبة %","قيمة الضريبة","الإجمالي"])
    ws.append(["دجاج 600 جم",10,"كرتون",120,15,180,1380])
    b=io.BytesIO(); wb.save(b); return b.getvalue()

def test_excel_engine_detects_header_and_validates():
    data=make_xlsx()
    lines=app.parse_xlsx_bytes(data)
    meta=app.analyze_xlsx_structure(data)
    assert len(lines)==1
    assert meta["primary"]["header_row"]==3
    assert meta["primary"]["column_map"]["اسم الصنف"]==1
    assert lines[0]["extraction_status"]=="VERIFIED"
    assert lines[0]["extraction_confidence"]>=90

def make_realistic_bilingual_invoice():
    wb=Workbook(); ws=wb.active; ws.title='فاتورة_ثمار_الغدا'
    ws.append(['شركة ثمار الغدا للمواد الغذائية - Thomar Al-Ghadda Foodstuff Company'])
    ws.append(['فاتورة ضريبية | TAX INVOICE'])
    ws.append([])
    ws.append(['اسم المورد / Supplier:', 'شركة ثمار الغدا للمواد الغذائية', None, None, 'رقم الفاتورة / Voucher No:', '2519458'])
    ws.append(['السجل التجاري / C.R.:', '4030543155', None, None, 'التاريخ / Date:', '20-Aug-2026'])
    ws.append(['الرقم الضريبي / VAT No:', '312040307300003'])
    ws.append(['اسم العميل / Customer:', 'مؤسسة منار تربة التجارية C1'])
    ws.append(['الرقم الضريبي للعميل / VAT:', '310333540400003'])
    ws.append([])
    ws.append(['م / S.No','رقم الصنف / Item No','الوصف / Description','الكمية / Qty','الوحدة / Unit','السعر / Price','المبلغ الخاضع للضريبة / Taxable','الضريبة (15%) / VAT','المجموع الشامل / Total'])
    ws.append([1,'1131','POTATO HYFARM 2.5KG X 4PKT\nبطاطس هيفارم 2.5 كجم × 4 عبوات',120,'CTN',50,None,None,None])
    ws.append([2,'1226','EMPEROR PANGASIUS STEAK 1KG X 10PKT\nشريحة لحم سمك الإمبراطور 10 كجم',5,'CTN',97,None,None,None])
    ws.append([3,'1175','COMMON CORP-200/300 CLEAN\nسمك كارب نظيف 200/300',5,'CTN',100,None,None,None])
    ws.append([4,'1213','POMFRET FISH 400/600-10KG\nسمك بومفريت غير نظيف',5,'CTN',125,None,None,None])
    ws.append([5,'1567','GUAVA PULP PREMIER 1KG X 16PKT\nجوافة بريمير',1,'CTN',80,None,None,None])
    for i in range(6,24):
        ws.append([i,str(1000+i),f'صنف تجريبي {i}',i%7+1,'CTN',20+i,None,None,None])
    ws.append([])
    ws.append([None,None,None,None,None,'الإجمالي غير شامل ضريبة القيمة المضافة:',None,None,None])
    ws.append([None,None,None,None,None,'مجموع الخصومات:',None,None,0])
    ws.append([None,None,None,None,None,'الإجمالي الخاضع للضريبة:',None,None,None])
    ws.append([None,None,None,None,None,'مجموع ضريبة القيمة المضافة (15%):',None,None,None])
    ws.append([None,None,None,None,None,'المبلغ الصافي شامل الضريبة:',None,None,None])
    ws.append([None,None,None,None,None,'بيانات البنك:','البنك الأهلي السعودي (SNB)'])
    b=io.BytesIO(); wb.save(b); return b.getvalue()


def test_real_bilingual_invoice_extracts_correct_columns():
    data=make_realistic_bilingual_invoice()
    meta=app.analyze_xlsx_structure(data)
    assert meta['primary']['header_row']==10
    assert meta['primary']['column_map']['كود المورد']==2
    assert meta['primary']['column_map']['اسم الصنف']==3
    assert meta['primary']['column_map']['الكمية']==4
    assert meta['primary']['column_map']['الوحدة']==5
    assert meta['primary']['column_map']['سعر الوحدة']==6
    assert meta['primary']['column_map']['الصافي']==7
    assert meta['primary']['column_map']['نسبة الضريبة']==8
    assert meta['primary']['column_map']['الإجمالي']==9
    assert meta['primary']['invoice']['invoice_number']=='2519458'
    lines=app.parse_xlsx_bytes(data)
    assert len(lines)==23
    first=lines[0]
    assert first['supplier_item_code']=='1131'
    assert first['qty']==120
    assert first['unit']=='CTN'
    assert first['unit_price']==50
    assert first['vat_rate']==15
    assert round(first['net_amount'],2)==6000
    assert round(first['gross_amount'],2)==6900
    assert all(x['extraction_status']=='VERIFIED' for x in lines)


def test_footer_rows_are_not_items():
    data=make_realistic_bilingual_invoice()
    lines=app.parse_xlsx_bytes(data)
    assert not any('البنك الأهلي' in (x.get('supplier_item_code') or '') for x in lines)


def test_unknown_header_does_not_guess_numeric_columns():
    wb=Workbook(); ws=wb.active
    ws.append(['X','Y','Z','A'])
    ws.append(['abc',10,120,999])
    b=io.BytesIO(); wb.save(b)
    meta=app.analyze_xlsx_structure(b.getvalue())
    assert meta['primary'] is None
